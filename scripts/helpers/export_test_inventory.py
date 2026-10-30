#!/usr/bin/env python3
"""Export the executable L1/L2/L3 migration inventory from pytest collection."""

from __future__ import annotations

import argparse
import re
import subprocess
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

NODE_RE = re.compile(r"^tests/.+\.py::")
MODEL_RE = re.compile(r"(?:model_id|model_path)\s*=\s*[furbFURB]*['\"]([^'\"]+)['\"]")


def collect_nodes(python: str) -> list[str]:
    proc = subprocess.run(
        [python, "-m", "pytest", "tests", "-o", "addopts=", "--collect-only", "-q"],
        check=True,
        text=True,
        stdout=subprocess.PIPE,
    )
    return [line.strip() for line in proc.stdout.splitlines() if NODE_RE.match(line.strip())]


def classify(path: str) -> tuple[str, str, str, str]:
    parts = Path(path).parts
    if "/model/" in path or path.startswith("tests/benchmark/models/"):
        return "L1", "保持", "真实模型场景或模型级 benchmark", "已迁移"
    if "/module/" in path:
        return "L2", "保持", "跨组件机制，受 L2 无模型 AST 守卫约束", "已迁移"
    if "/unit/" in path or path.startswith("tests/helpers/tests/") or path.startswith("tests/regression/scripts/"):
        return "L3", "保持", "纯函数、parser、validator、默认值或异常契约", "已迁移"
    if path.startswith("tests/smoke/"):
        return "L1", "保持", "真实小模型 smoke 场景", "边界保留"
    if path.startswith("tests/benchmark/"):
        return "L2", "保持", "算子/微基准机制", "边界保留"
    if path.startswith("tests/regression/"):
        component = parts[2] if len(parts) > 2 else "regression"
        return "范围外", "维护边界", f"{component} 既有套件；新增用例须按三问判定树路由", "边界登记"
    return "范围外", "维护边界", "非 regression 测试资产", "边界登记"


def schedule(path: str, source: str) -> str:
    if path.startswith("tests/benchmark/"):
        return "benchmark/nightly"
    if "pytest.mark.network" in source or "pytest.mark.nightly" in source:
        return "nightly"
    return "incremental"


def module_name(path: str) -> str:
    parts = Path(path).parts
    if path.startswith("tests/regression/") and len(parts) > 2:
        return parts[2]
    if path.startswith("tests/benchmark/") and len(parts) > 2:
        return f"benchmark/{parts[2]}"
    return parts[1] if len(parts) > 1 else "tests"


def source_facts(root: Path, path: str) -> tuple[str, str]:
    file_path = root / path
    source = file_path.read_text(encoding="utf-8") if file_path.is_file() else ""
    identities = sorted(set(MODEL_RE.findall(source)))
    return source, ", ".join(identities) if identities else "无显式模型身份"


def style_sheet(sheet) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, column in enumerate(sheet.columns, start=1):
        values = [str(cell.value or "") for cell in list(column)[:200]]
        sheet.column_dimensions[get_column_letter(index)].width = min(max(max(map(len, values), default=8) + 2, 10), 55)


def write_workbook(rows: list[dict[str, str]], output: Path, commit: str) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    summary = workbook.create_sheet("总览")
    summary.append(["指标", "值", "说明"])
    counts = Counter(row["目标责任层"] for row in rows)
    schedules = Counter(row["CI调度"] for row in rows)
    summary_rows = [
        ("built_from_commit", commit, "生成时仓库 HEAD"),
        ("collect_total", len(rows), "pytest tests -o addopts= --collect-only -q"),
        *[(f"目标层/{name}", counts[name], "逐 node 计数") for name in ("L1", "L2", "L3", "范围外")],
        *[(f"调度/{name}", count, "逐 node 计数") for name, count in sorted(schedules.items())],
    ]
    for item in summary_rows:
        summary.append(item)

    headers = list(rows[0]) if rows else []
    all_cases = workbook.create_sheet("全部用例")
    all_cases.append(headers)
    for row in rows:
        all_cases.append([row[key] for key in headers])

    for title, layer in (("L1模型级", "L1"), ("L2模块级", "L2"), ("L3单元级", "L3")):
        sheet = workbook.create_sheet(title)
        sheet.append(headers)
        for row in rows:
            if row["目标责任层"] == layer:
                sheet.append([row[key] for key in headers])

    review = workbook.create_sheet("迁移与复核")
    review_headers = ["当前node_id", "目标node_id", "迁移动作", "状态", "test_map要求", "断言/sole-watcher复核"]
    review.append(review_headers)
    for row in rows:
        review.append(
            [
                row["node_id"],
                row["node_id"],
                row["迁移动作"],
                row["复核状态"],
                "路径变化后刷新 built_from_commit；old=0/new可collect",
                "删除或合并前核对 assertion_manifest 与 l1_watchers",
            ]
        )

    for sheet in workbook.worksheets:
        style_sheet(sheet)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--python", default="python")
    parser.add_argument("--output", type=Path, default=Path("docs/design/l1_l2_l3_test_inventory.xlsx"))
    args = parser.parse_args()
    root = Path.cwd()
    commit = subprocess.run(["git", "rev-parse", "HEAD"], check=True, text=True, stdout=subprocess.PIPE).stdout.strip()
    rows: list[dict[str, str]] = []
    cache: dict[str, tuple[str, str]] = {}
    for node_id in collect_nodes(args.python):
        path = node_id.split("::", 1)[0]
        source, identity = cache.setdefault(path, source_facts(root, path))
        layer, action, evidence, review = classify(path)
        rows.append(
            {
                "node_id": node_id,
                "当前路径": path,
                "当前责任层": layer,
                "目标责任层": layer,
                "目标模块": module_name(path),
                "模型身份": identity,
                "构建物": "L1RunArtifact" if layer == "L1" else ("canonical synthetic artifact" if layer == "L2" else "无模型构建"),
                "marker": "nightly/network" if schedule(path, source) == "nightly" else "default",
                "CI调度": schedule(path, source),
                "迁移PR": "PR2/PR3" if layer == "L1" else ("PR4/PR5" if layer in {"L2", "L3"} else "维护边界"),
                "迁移动作": action,
                "分类证据": evidence,
                "置信度": "高" if layer != "范围外" else "边界",
                "复核状态": review,
            }
        )
    write_workbook(rows, args.output, commit)
    print(f"wrote {args.output} with {len(rows)} collected nodes")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
